import tkinter as tk
from tkinter import ttk, messagebox
from openpyxl import Workbook, load_workbook
import os


# -------------------------------------------------
# Excel File
# -------------------------------------------------

FILE_NAME = "student_results.xlsx"

COLUMNS = [
    "Name",
    "Roll No.",
    "Class",
    "Subject 1",
    "Subject 2",
    "Subject 3",
    "Subject 4",
    "Subject 5",
    "Total Marks",
    "Percentage",
    "Result"
]


def create_excel_file():
    """Create Excel file if it does not already exist."""
    if not os.path.exists(FILE_NAME):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Results"

        for col, heading in enumerate(COLUMNS, start=1):
            sheet.cell(row=1, column=col, value=heading)

        workbook.save(FILE_NAME)


# -------------------------------------------------
# Save Student
# -------------------------------------------------

def save_student():
    name = name_entry.get().strip()
    roll_no = roll_entry.get().strip()
    student_class = class_entry.get().strip()

    marks_entries = [
        sub1_entry,
        sub2_entry,
        sub3_entry,
        sub4_entry,
        sub5_entry
    ]

    # Check empty fields
    if not name or not roll_no or not student_class:
        messagebox.showerror(
            "Error",
            "Please enter Name, Roll No. and Class."
        )
        return

    # Check roll number
    if not roll_no.isdigit():
        messagebox.showerror(
            "Error",
            "Roll No. must contain numbers only."
        )
        return

    # Read marks
    marks = []

    for entry in marks_entries:
        value = entry.get().strip()

        if not value.isdigit():
            messagebox.showerror(
                "Error",
                "All subject marks must be numbers."
            )
            return

        mark = int(value)

        if mark < 0 or mark > 100:
            messagebox.showerror(
                "Error",
                "Marks must be between 0 and 100."
            )
            return

        marks.append(mark)

    # Convert roll number
    roll_number = int(roll_no)

    # Check duplicate roll number
    workbook = load_workbook(FILE_NAME)
    sheet = workbook["Results"]

    for row in sheet.iter_rows(min_row=2, values_only=True):
        if row[1] == roll_number:
            messagebox.showerror(
                "Error",
                "A student with this Roll No. already exists."
            )
            workbook.close()
            return

    # Calculate total
    total = sum(marks)

    # Calculate percentage
    percentage = total / 5

    # Result
    if percentage >= 40:
        result = "Pass"
    else:
        result = "Fail"

    # Save record
    sheet.append([
        name,
        roll_number,
        student_class,
        marks[0],
        marks[1],
        marks[2],
        marks[3],
        marks[4],
        total,
        percentage,
        result
    ])

    workbook.save(FILE_NAME)
    workbook.close()

    messagebox.showinfo(
        "Success",
        f"Student saved successfully!\n\n"
        f"Total Marks: {total}\n"
        f"Percentage: {percentage:.2f}%\n"
        f"Result: {result}"
    )

    clear_fields()


# -------------------------------------------------
# Clear Fields
# -------------------------------------------------

def clear_fields():
    entries = [
        name_entry,
        roll_entry,
        class_entry,
        sub1_entry,
        sub2_entry,
        sub3_entry,
        sub4_entry,
        sub5_entry
    ]

    for entry in entries:
        entry.delete(0, tk.END)

    roll_result_entry.delete(0, tk.END)

    # result_label.config(text="")
    

# -------------------------------------------------
# Get Result
# -------------------------------------------------

def get_result():
    roll_no = roll_result_entry.get().strip()

    if not roll_no:
        messagebox.showerror(
            "Error",
            "Please enter Roll No."
        )
        return

    if not roll_no.isdigit():
        messagebox.showerror(
            "Error",
            "Roll No. must contain numbers only."
        )
        return

    roll_number = int(roll_no)

    workbook = load_workbook(FILE_NAME)
    sheet = workbook["Results"]

    found = False

    for row in sheet.iter_rows(min_row=2, values_only=True):

        if row[1] == roll_number:
            found = True

            name = row[0]
            student_class = row[2]
            total = row[8]
            percentage = row[9]
            result = row[10]

            # Clear previous result
            for item in result_tree.get_children():
                result_tree.delete(item)

            result_tree.insert(
                "",
                tk.END,
                values=(
                    name,
                    row[1],
                    student_class,
                    total,
                    f"{percentage:.2f}%",
                    result
                )
            )

            break

    workbook.close()

    if not found:
        messagebox.showerror(
            "Not Found",
            "❌ Student record not found."
        )


# -------------------------------------------------
# Main Window
# -------------------------------------------------

create_excel_file()

window = tk.Tk()

window.title("Student Result Management System")
window.geometry("1000x700")
window.resizable(False, False)


# -------------------------------------------------
# Title
# -------------------------------------------------

title_label = tk.Label(
    window,
    text="STUDENT RESULT MANAGEMENT SYSTEM",
    font=("Arial", 22, "bold")
)

title_label.pack(pady=15)


# -------------------------------------------------
# Notebook / Tabs
# -------------------------------------------------

notebook = ttk.Notebook(window)
notebook.pack(fill="both", expand=True, padx=20, pady=10)


# =================================================
# ADD STUDENT TAB
# =================================================

add_frame = ttk.Frame(notebook)

notebook.add(
    add_frame,
    text="Add Student"
)


# Student details heading

heading = tk.Label(
    add_frame,
    text="Student Details",
    font=("Arial", 16, "bold")
)

heading.grid(
    row=0,
    column=0,
    columnspan=2,
    pady=15
)


# Name

tk.Label(
    add_frame,
    text="Name:",
    font=("Arial", 12)
).grid(
    row=1,
    column=0,
    padx=20,
    pady=8,
    sticky="e"
)

name_entry = tk.Entry(
    add_frame,
    width=30,
    font=("Arial", 12)
)

name_entry.grid(
    row=1,
    column=1,
    pady=8
)


# Roll No

tk.Label(
    add_frame,
    text="Roll No.:",
    font=("Arial", 12)
).grid(
    row=2,
    column=0,
    padx=20,
    pady=8,
    sticky="e"
)

roll_entry = tk.Entry(
    add_frame,
    width=30,
    font=("Arial", 12)
)

roll_entry.grid(
    row=2,
    column=1,
    pady=8
)


# Class

tk.Label(
    add_frame,
    text="Class:",
    font=("Arial", 12)
).grid(
    row=3,
    column=0,
    padx=20,
    pady=8,
    sticky="e"
)

class_entry = tk.Entry(
    add_frame,
    width=30,
    font=("Arial", 12)
)

class_entry.grid(
    row=3,
    column=1,
    pady=8
)


# Subject marks

subject_entries = []

subjects = [
    "Subject 1:",
    "Subject 2:",
    "Subject 3:",
    "Subject 4:",
    "Subject 5:"
]

for i, subject in enumerate(subjects, start=4):

    tk.Label(
        add_frame,
        text=subject,
        font=("Arial", 12)
    ).grid(
        row=i,
        column=0,
        padx=20,
        pady=6,
        sticky="e"
    )

    entry = tk.Entry(
        add_frame,
        width=30,
        font=("Arial", 12)
    )

    entry.grid(
        row=i,
        column=1,
        pady=6
    )

    subject_entries.append(entry)


sub1_entry = subject_entries[0]
sub2_entry = subject_entries[1]
sub3_entry = subject_entries[2]
sub4_entry = subject_entries[3]
sub5_entry = subject_entries[4]


# Save button

save_button = tk.Button(
    add_frame,
    text="💾 Save",
    font=("Arial", 12, "bold"),
    width=15,
    command=save_student
)

save_button.grid(
    row=9,
    column=0,
    columnspan=2,
    pady=20
)


# =================================================
# GET RESULT TAB
# =================================================

get_frame = ttk.Frame(notebook)

notebook.add(
    get_frame,
    text="Get Result"
)


tk.Label(
    get_frame,
    text="Enter Roll No.:",
    font=("Arial", 14)
).pack(pady=(40, 10))


roll_result_entry = tk.Entry(
    get_frame,
    width=30,
    font=("Arial", 13)
)

roll_result_entry.pack()


get_button = tk.Button(
    get_frame,
    text="🔍 Get Result",
    font=("Arial", 12, "bold"),
    width=15,
    command=get_result
)

get_button.pack(pady=15)


# Result Treeview

result_columns = (
    "Name",
    "Roll No.",
    "Class",
    "Total Marks",
    "Percentage",
    "Result"
)

result_tree = ttk.Treeview(
    get_frame,
    columns=result_columns,
    show="headings",
    height=8
)

for column in result_columns:

    result_tree.heading(
        column,
        text=column
    )

    result_tree.column(
        column,
        width=130,
        anchor="center"
    )


result_tree.pack(
    fill="x",
    padx=30,
    pady=20
)


# =================================================
# SHOW ALL RESULTS TAB
# =================================================

all_frame = ttk.Frame(notebook)

notebook.add(
    all_frame,
    text="Show All Results"
)


# show_button = tk.Button(
#     all_frame,
#     text="📋 Show All Results",
#     font=("Arial", 12, "bold"),
#     command=show_all_results
# )

# show_button.pack(pady=20)


# All results Treeview

all_columns = (
    "Name",
    "Roll No.",
    "Class",
    "Total Marks",
    "Percentage",
    "Result"
)

all_tree = ttk.Treeview(
    all_frame,
    columns=all_columns,
    show="headings",
    height=18
)

for column in all_columns:

    all_tree.heading(
        column,
        text=column
    )

    all_tree.column(
        column,
        width=140,
        anchor="center"
    )

all_tree.pack(
    fill="both",
    expand=True,
    padx=20,
    pady=10
)


# -------------------------------------------------
# Show All Results
# -------------------------------------------------

def show_all_results():
    # Clear existing rows from the Show All Results table
    for item in all_tree.get_children():
        all_tree.delete(item)

    try:
        workbook = load_workbook(FILE_NAME)
        sheet = workbook["Results"]

        if sheet.max_row <= 1:
            workbook.close()
            messagebox.showinfo("No Records", "No student records found.")
            return

        for row in sheet.iter_rows(min_row=2, values_only=True):
            if not row[0]:
                continue

            percentage = row[9] if row[9] is not None else 0

            all_tree.insert(
                "",
                tk.END,
                values=(
                    row[0],
                    row[1],
                    row[2],
                    row[8],
                    f"{float(percentage):.2f}%",
                    row[10]
                )
            )

        workbook.close()

    except FileNotFoundError:
        messagebox.showerror("Error", "student_results.xlsx file not found.")

    except Exception as e:
        messagebox.showerror("Error", f"Unable to read Excel file.\n\n{e}")




# -------------------------------------------------
# Modified Show All function for the second Treeview
# -------------------------------------------------

def show_all_results():

    # Delete old records
    for item in all_tree.get_children():
        all_tree.delete(item)

    workbook = load_workbook(FILE_NAME)
    sheet = workbook["Results"]

    for row in sheet.iter_rows(min_row=2, values_only=True):

        percentage = row[9]

        all_tree.insert(
            "",
            tk.END,
            values=(
                row[0],
                row[1],
                row[2],
                row[8],
                f"{percentage:.2f}%",
                row[10]
            )
        )

    workbook.close()

show_button = tk.Button(
    all_frame,
    text="📋 Show All Results",
    font=("Arial", 12, "bold"),
    command=show_all_results
)

show_button.pack(pady=20)

# -------------------------------------------------
# Run Application
# -------------------------------------------------

window.mainloop()